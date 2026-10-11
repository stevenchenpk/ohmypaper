"""
生成格式体检的回归样例（虚构内容，不含真实论文）。

    python3 tests/fixtures/make_samples.py

需要 python-docx（pip install python-docx）。生成的 docx 已提交到仓库，
平时跑测试不需要 Python；改了样例才需要重新生成。

- good.docx：完全符合 tests/fixtures/test-profile.yaml 的论文
- bad.docx：同样的论文，故意埋了一批常见格式错误，每个错误都对应测试里的一个断言
"""

import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(__file__).parent

REFS = [
    "LECUN Y, BENGIO Y, HINTON G. Deep learning[J]. Nature, 2015, 521(7553): 436-444.",
    "张三, 李四. 生成式人工智能对高等教育的影响研究[J]. 中国高教研究, 2024(3): 12-19.",
    "王五. 某某问题研究[D]. 北京: 北京大学, 2020.",
    "教育部. 关于某事项的通知[EB/OL]. (2025-03-01)[2026-10-01]. http://www.moe.gov.cn/x.html.",
]


def set_fonts(run_or_style, cn, en, size=None, bold=None):
    font = run_or_style.font
    font.name = en
    rpr = run_or_style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for k in ("ascii", "hAnsi"):
        rfonts.set(qn(f"w:{k}"), en)
    rfonts.set(qn("w:eastAsia"), cn)
    for k in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme"):
        rfonts.attrib.pop(qn(f"w:{k}"), None)
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold


def page_field(paragraph):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def page_numbering(section, fmt=None, start=None):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sectPr.append(pg)
    if fmt:
        pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))


def setup_footer(section, header_text, with_page=True):
    section.footer.is_linked_to_previous = False
    section.header.is_linked_to_previous = False
    fp = section.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if with_page:
        page_field(fp)
    hp = section.header.paragraphs[0]
    hp.text = header_text
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def margins(section, top=2.5, bottom=2.5, left=3.0, right=2.5):
    section.top_margin, section.bottom_margin = Cm(top), Cm(bottom)
    section.left_margin, section.right_margin = Cm(left), Cm(right)
    section.page_width, section.page_height = Cm(21), Cm(29.7)


def build(bad: bool):
    doc = Document()
    core = doc.core_properties
    core.author = "张同学" if bad else ""
    core.last_modified_by = "张同学" if bad else ""
    core.created = core.modified = datetime.datetime(2026, 3, 1)

    # 样式：正文 宋体/TNR 小四 1.5 倍 首行缩进 2 字符 两端对齐
    normal = doc.styles["Normal"]
    if bad:
        set_fonts(normal, "等线", "Calibri", 10.5)  # 错：字体字号
    else:
        set_fonts(normal, "宋体", "Times New Roman", 12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE if bad else WD_LINE_SPACING.ONE_POINT_FIVE  # 错：行距
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = pf.space_after = Pt(0)
    if not bad:
        ind = OxmlElement("w:ind")
        ind.set(qn("w:firstLineChars"), "200")
        ind.set(qn("w:firstLine"), "480")
        normal.element.get_or_add_pPr().append(ind)

    for level, size in ((1, 16), (2, 14), (3, 12)):
        st = doc.styles[f"Heading {level}"]
        set_fonts(st, "黑体", "Times New Roman", size, bold=True)
        st.font.color.rgb = None
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.first_line_indent = Pt(0)

    # 第 1 节：封面
    sec = doc.sections[0]
    margins(sec)
    doc.add_paragraph("某某大学本科毕业论文").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("论文题目：示例论文").alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 第 2 节：摘要、目录（罗马数字）
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    margins(sec)
    setup_footer(sec, "某某大学本科毕业论文")
    page_numbering(sec, "upperRoman", 1)
    doc.add_heading("摘  要", level=1)
    doc.add_paragraph("本文以一个虚构的问题为例，说明论文格式体检工具如何工作。这段文字只用于测试。")
    doc.add_heading("目  录", level=1)
    doc.add_paragraph("第一章 绪论 ........ 1")

    # 第 3 节：正文（阿拉伯数字，从 1 开始）
    if bad:
        # 错：正文和摘要目录不分节
        pass
    else:
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        margins(sec)
        setup_footer(sec, "某某大学本科毕业论文")
        page_numbering(sec, "decimal", 1)

    doc.add_heading("第一章 绪论", level=1)
    doc.add_heading("1.1 研究背景", level=2)
    doc.add_paragraph("深度学习已经改变了很多研究领域[1]，生成式人工智能也开始影响高等教育[2]。本节说明研究背景。")
    doc.add_heading("1.2 研究意义", level=2)
    p = doc.add_paragraph("关于这个问题已有学位论文做过讨论[3]，政策层面也有相关通知[4]。")
    if bad:
        p.add_run("如图 3-1 所示，错误！未找到引用源。")  # 错：交叉引用失效
        doc.add_paragraph("此外，另一项研究也值得参考[6]。")  # 错：引用不存在的 [6]

    if bad:
        doc.add_heading("第三章 研究方法", level=1)  # 错：第一章后面直接第三章
        doc.add_heading("3.1.1 数据来源", level=3)  # 错：跳级
        doc.add_paragraph("3.2 研究设计")  # 错：疑似标题未用样式（下面加粗）
        doc.paragraphs[-1].runs[0].bold = True
    else:
        doc.add_heading("第二章 研究方法", level=1)
        doc.add_heading("2.1 数据来源", level=2)
        doc.add_heading("2.1.1 样本", level=3)
    body = doc.add_paragraph("本章介绍研究所用的数据和方法，所有数据均为虚构。")
    if bad:
        set_fonts(body.runs[0], "楷体", "Arial")  # 错：直接格式改了字体

    doc.add_heading("参考文献", level=1)
    refs = list(REFS)
    if bad:
        refs[0] = refs[0].replace("[J]", "")  # 错：缺文献类型标识
        refs[3] = "教育部. 关于某事项的通知[EB]. http://www.moe.gov.cn/x.html."  # 错：缺 /OL 和引用日期
        refs.append("赵六. 一篇没有被引用的文章[J]. 某某学报, 2023(1): 1-5.")
    for i, r in enumerate(refs, 1):
        text = r if (bad and i == 2) else f"[{i}] {r}"  # 错：第 2 条漏了序号
        rp = doc.add_paragraph(text)
        rp.paragraph_format.first_line_indent = Pt(0)
        for run in rp.runs:
            set_fonts(run, "宋体", "Times New Roman", 10.5)

    doc.add_heading("致  谢", level=1)
    doc.add_paragraph("感谢所有帮助过我的人。")

    sec = doc.sections[-1]
    if bad:
        margins(sec, left=2.5)  # 错：左边距
        # 错：正文所在的节没有页码，也没有页眉
        sec.footer.is_linked_to_previous = False
        sec.footer.paragraphs[0].text = ""
        sec.header.is_linked_to_previous = False
        sec.header.paragraphs[0].text = ""

    return doc


if __name__ == "__main__":
    build(False).save(OUT / "good.docx")
    build(True).save(OUT / "bad.docx")
    print("wrote good.docx, bad.docx")
