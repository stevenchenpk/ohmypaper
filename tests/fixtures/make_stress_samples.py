"""
生成"更像真论文"的合格样例，专门用来压测格式体检的误报（虚构内容）。

    python3 tests/fixtures/make_stress_samples.py

需要 python-docx；captions-libreoffice.docx 还需要 LibreOffice（soffice）。
所有样例都应当完全符合 tests/fixtures/test-profile.yaml，输出到 tests/fixtures/stress/：

- toc-field.docx：Word 真实的目录域（TOC 域 + toc 1/toc 2 样式 + PAGEREF + 超链接 + 书签）
- landscape.docx：正文中间插入横向节（放宽表格），再回到纵向；页码续前节，页眉页脚链接到前一节
- captions.docx：插图（内嵌图片）+ 题注样式 + SEQ 域、表格上方的表题、居中公式行、制表位排版的公式编号行、交叉引用 REF 域
- mixed-numbering.docx：标题用 Word 多级列表自动编号（标题文字里没有编号），正文混有手写编号段落和自动编号列表，参考文献用 "[%1]" 自动编号
- wps-like.docx：模仿 WPS 的写法：样式 id 是数字，字体字号行距缩进全部是直接格式，rFonts 带 hint="eastAsia"，页脚页码放在文本框里，app.xml 写 WPS Office
- captions-libreoffice.docx：captions.docx 经 LibreOffice 转 odt 再转回 docx
"""

import datetime
import io
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import zipfile
import zlib
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT, WD_SECTION  # noqa: E402
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER  # noqa: E402
from docx.oxml import OxmlElement, parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

from make_samples import REFS, margins, page_field, page_numbering, set_fonts, setup_footer  # noqa: E402

OUT = Path(__file__).parent / "stress"
HEADER = "某某大学本科毕业论文"

ALL_REFS = REFS + [
    "Vaswani A, Shazeer N, Parmar N, et al. Attention is all you need[C]//Proceedings of the 31st International "
    "Conference on Neural Information Processing Systems. Long Beach: Curran Associates, 2017: 6000-6010.",
]

ABSTRACT_CN = "本文以一个虚构的问题为例，说明论文格式体检工具如何工作。全文所有数据均为虚构，只用于测试格式检查是否会误报。"
ABSTRACT_EN = (
    "This thesis uses a fictional problem to show how the format checker works. "
    "All data are fictional and only used for testing false positives."
)

# 正文内容：(kind, text)；kind = h1/h2/h3/p
CONTENT = [
    ("h1", "第一章 绪论"),
    ("h2", "1.1 研究背景"),
    ("p", "近年来，深度学习（Deep Learning）在计算机视觉、自然语言处理等领域取得了显著进展[1]。"
          "生成式人工智能（Generative AI）的快速发展也开始深刻影响高等教育的教学与评价方式[2]。"),
    ("p", "与此同时，学位论文的格式规范日益受到重视。以往研究多关注论文的内容质量[3]，"
          "对格式检查工具的讨论较少，本文尝试补充这一方面的工作。"),
    ("h2", "1.2 研究意义"),
    ("p", "本研究的意义主要体现在两个方面：一是为学生提供便捷的自查工具，二是减轻指导教师的重复劳动。"
          "教育主管部门近年也发布了加强学位论文管理的相关通知[4]。"),
    ("h2", "1.3 论文结构"),
    ("p", "全文共分为四章。第一章为绪论；第二章介绍相关技术；第三章给出系统设计；第四章总结全文并展望后续工作。"),
    ("h1", "第二章 相关技术"),
    ("h2", "2.1 Transformer 模型"),
    ("p", "Transformer 模型完全基于注意力机制（Attention Mechanism），摒弃了循环和卷积结构[5]。"
          "它在机器翻译任务上取得了当时最好的效果，此后成为大语言模型的基础架构。"),
    ("h3", "2.1.1 自注意力机制"),
    ("p", "自注意力机制允许序列中的每个位置关注其他所有位置，从而捕获长距离依赖关系，并且便于并行计算。"),
    ("h3", "2.1.2 位置编码"),
    ("p", "由于模型本身不包含顺序信息，需要额外引入位置编码。常用的做法是正弦位置编码和可学习的位置向量。"),
    ("h2", "2.2 本章小结"),
    ("p", "本章回顾了与本文相关的主要技术，为后续的系统设计打下基础。"),
    ("h1", "第三章 系统设计"),
    ("h2", "3.1 总体架构"),
    ("p", "系统由文件解析、结构识别、规则检查和报告生成四个模块组成，各模块之间通过统一的数据结构交换信息。"),
    ("h2", "3.2 模块划分"),
    ("p", "文件解析模块负责读取 docx 文件，结构识别模块负责划分前置部分、正文和参考文献，规则检查模块逐项比对学校要求。"),
    ("h1", "第四章 总结与展望"),
    ("p", "本文完成了一个论文格式体检工具的原型，后续将补充更多学校的规则档案，并改进结构识别的准确率。"),
]


# ---------- 通用小工具 ----------

def tiny_png(w=24, h=16):
    """标准库生成一张灰色小 PNG。"""
    raw = b"".join(b"\x00" + bytes([200, 200, 200]) * w for _ in range(h))

    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def run_el(text=None, rpr=None, fld=None, instr=None, tab=False):
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(rpr)
    if fld:
        r.append(el("w:fldChar", fldCharType=fld))
    if instr is not None:
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        r.append(it)
    if tab:
        r.append(OxmlElement("w:tab"))
    if text is not None:
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
    return r


def field(parent, instr, result, rpr_factory=lambda: None):
    """复杂域：begin / instrText / separate / 结果 / end，追加到 parent（段落或超链接）。"""
    parent.append(run_el(rpr=rpr_factory(), fld="begin"))
    parent.append(run_el(rpr=rpr_factory(), instr=instr))
    parent.append(run_el(rpr=rpr_factory(), fld="separate"))
    parent.append(run_el(result, rpr=rpr_factory()))
    parent.append(run_el(rpr=rpr_factory(), fld="end"))


def drop_restart(section):
    """python-docx 的 add_section 会把上一节的 pgNumType（含 start）复制到新节，正文后续节要续前节。"""
    pg = section._sectPr.find(qn("w:pgNumType"))
    if pg is not None:
        pg.attrib.pop(qn("w:start"), None)


def setup_styles(doc):
    core = doc.core_properties
    core.author = ""
    core.last_modified_by = ""
    core.created = core.modified = datetime.datetime(2026, 3, 1)

    normal = doc.styles["Normal"]
    set_fonts(normal, "宋体", "Times New Roman", 12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = pf.space_after = Pt(0)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:firstLineChars"), "200")
    ind.set(qn("w:firstLine"), "480")
    normal.element.get_or_add_pPr().append(ind)

    for level, size in ((1, 16), (2, 14), (3, 12)):
        st = doc.styles[f"Heading {level}"]
        set_fonts(st, "黑体", "Times New Roman", size, bold=True)
        st.font.color.rgb = None
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.first_line_indent = Pt(0)

    cap = doc.styles["Caption"]
    set_fonts(cap, "宋体", "Times New Roman", 10.5, bold=False)
    cap.font.color.rgb = RGBColor(0, 0, 0)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Pt(0)


def cover_and_front(doc, toc=None):
    """封面一节 + 摘要目录一节（罗马页码）。toc(doc) 负责写目录内容。"""
    sec = doc.sections[0]
    margins(sec)
    for t in ("某某大学本科毕业论文", "论文题目：示例论文"):
        p = doc.add_paragraph(t)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    margins(sec)
    setup_footer(sec, HEADER)
    page_numbering(sec, "upperRoman", 1)
    doc.add_heading("摘  要", level=1)
    doc.add_paragraph(ABSTRACT_CN)
    p = doc.add_paragraph("关键词：格式检查；学位论文；误报")
    p.paragraph_format.first_line_indent = Pt(0)
    doc.add_heading("ABSTRACT", level=1)
    doc.add_paragraph(ABSTRACT_EN)
    doc.add_heading("目  录", level=1)
    if toc:
        toc(doc)
    else:
        for level, text in [(k, t) for k, t in CONTENT if k in ("h1", "h2")]:
            p = doc.add_paragraph(f"{text} ........ 1")
            p.paragraph_format.first_line_indent = Pt(0)


def body_section(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    margins(sec)
    setup_footer(sec, HEADER)
    page_numbering(sec, "decimal", 1)
    return sec


def add_content(doc, content, on_heading=None):
    for kind, text in content:
        if kind == "p":
            doc.add_paragraph(text)
        else:
            h = doc.add_heading(text, level=int(kind[1]))
            if on_heading:
                on_heading(h, kind, text)


def back_matter(doc, refs=ALL_REFS, numbered=None):
    doc.add_heading("参考文献", level=1)
    for i, r in enumerate(refs, 1):
        rp = doc.add_paragraph(r if numbered else f"[{i}] {r}")
        rp.paragraph_format.first_line_indent = Pt(0)
        if numbered:
            numbered(rp)
        for run in rp.runs:
            set_fonts(run, "宋体", "Times New Roman", 10.5)
    doc.add_heading("致  谢", level=1)
    doc.add_paragraph("感谢指导教师在论文写作过程中给予的帮助，也感谢同学们的支持。")


# ---------- 1. 目录域 ----------

def build_toc_field():
    doc = Document()
    setup_styles(doc)
    for name, left in (("toc 1", 0), ("toc 2", 420)):
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        st.next_paragraph_style = doc.styles["Normal"]
        st.paragraph_format.first_line_indent = Pt(0)
        st.paragraph_format.left_indent = Pt(left / 20)
        st.paragraph_format.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    entries = [(k, t) for k, t in CONTENT if k in ("h1", "h2")] + [("h1", "参考文献"), ("h1", "致  谢")]
    marks = {t: f"_Toc{180000000 + i}" for i, (_, t) in enumerate(entries)}

    def toc(doc):
        for i, (kind, text) in enumerate(entries):
            p = doc.add_paragraph(style=f"toc {kind[1]}")
            if i == 0:
                p._p.append(run_el(fld="begin"))
                p._p.append(run_el(instr=' TOC \\o "1-3" \\h \\z \\u '))
                p._p.append(run_el(fld="separate"))
            link = OxmlElement("w:hyperlink")
            link.set(qn("w:anchor"), marks[text])
            link.set(qn("w:history"), "1")
            link.append(run_el(text))
            link.append(run_el(tab=True))
            web_hidden = lambda: parse_xml(f"<w:rPr {nsdecls('w')}><w:webHidden/></w:rPr>")  # noqa: E731
            field(link, f" PAGEREF {marks[text]} \\h ", str(1 + i // 3), web_hidden)
            p._p.append(link)
        end = doc.add_paragraph()
        end._p.append(run_el(fld="end"))

    cover_and_front(doc, toc)
    body_section(doc)

    bm_id = [0]

    def bookmark(h, kind, text):
        if text not in marks:
            return
        bm_id[0] += 1
        start = el("w:bookmarkStart", id=bm_id[0], name=marks[text])
        h._p.insert(1, start)  # pPr 之后
        h._p.append(el("w:bookmarkEnd", id=bm_id[0]))

    add_content(doc, CONTENT, bookmark)
    back_matter(doc)
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") and p.text in marks and not p._p.findall(qn("w:bookmarkStart")):
            bookmark(p, None, p.text)
    return doc


# ---------- 2. 横向节 ----------

def build_landscape():
    doc = Document()
    setup_styles(doc)
    cover_and_front(doc)
    body_section(doc)
    split = CONTENT.index(("h2", "2.2 本章小结"))
    add_content(doc, CONTENT[:split])
    doc.add_paragraph("各模型在公开数据集上的详细结果见表 2-1，由于列数较多，该表放在横向页面中。")

    # 横向节：A4 横放，边距数值不变；页眉页脚链接到前一节，页码续前节
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    drop_restart(sec)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.5)
    cap = doc.add_paragraph("表 2-1 各模型在测试集上的结果", style="Caption")
    cols = ["模型", "参数量", "BLEU", "ROUGE-L", "准确率", "召回率", "F1", "推理时间（ms）", "显存（GB）"]
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = doc.styles["Table Grid"]
    for c, name in zip(table.rows[0].cells, cols):
        c.text = name
    for row in [
        ["Baseline", "110M", "27.3", "41.2", "0.812", "0.790", "0.801", "35", "2.1"],
        ["Transformer", "213M", "28.4", "43.5", "0.845", "0.821", "0.833", "42", "3.6"],
        ["本文方法", "220M", "29.1", "44.8", "0.861", "0.840", "0.850", "44", "3.8"],
    ]:
        for c, v in zip(table.add_row().cells, row):
            c.text = v
    for row in table.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    assert cap.text

    # 回到纵向
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    drop_restart(sec)
    sec.orientation = WD_ORIENT.PORTRAIT
    margins(sec)
    add_content(doc, CONTENT[split:])
    back_matter(doc)
    return doc


# ---------- 3. 图表题注 ----------

def caption(doc, label, chapter, seq_result, title, simple=False, bookmark=None):
    """题注：标签 + 章号 + SEQ 域 + 题名。交叉引用的书签像 Word 一样只包住"标签和编号"。"""
    p = doc.add_paragraph(style="Caption")
    if bookmark:
        bid = 900 + len(bookmark) + seq_result * 10 + chapter * 100
        p._p.append(el("w:bookmarkStart", id=bid, name=bookmark))
    p._p.append(run_el(f"{label} {chapter}-"))
    instr = f" SEQ {label} \\* ARABIC \\s 1 "
    if simple:
        fs = OxmlElement("w:fldSimple")
        fs.set(qn("w:instr"), instr)
        fs.append(run_el(str(seq_result)))
        p._p.append(fs)
    else:
        field(p._p, instr, str(seq_result))
    if bookmark:
        p._p.append(el("w:bookmarkEnd", id=bid))
    p._p.append(run_el(f" {title}"))
    return p


def ref_para(doc, before, ref_text, bookmark, after):
    p = doc.add_paragraph(before)
    field(p._p, f" REF {bookmark} \\h ", ref_text)
    p._p.append(run_el(after))
    return p


def figure(doc, chapter, n, title, ref):
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.first_line_indent = Pt(0)
    pic.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pic.add_run().add_picture(io.BytesIO(tiny_png()), width=Cm(8))
    caption(doc, "图", chapter, n, title, simple=(n == 2), bookmark=ref)


def small_table(doc, chapter, n, title, rows):
    caption(doc, "表", chapter, n, title)
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = doc.styles["Table Grid"]
    for r in rows:
        for c, v in zip(t.add_row().cells, r):
            c.text = v
            c.paragraphs[0].paragraph_format.first_line_indent = Pt(0)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def equation(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    return p


def tab_equation(doc, formula, number):
    """论文模板常见的公式行：中间制表位放公式，右侧制表位放编号，段落本身两端对齐、不缩进。"""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    ts = p.paragraph_format.tab_stops
    ts.add_tab_stop(Cm(7.75), WD_TAB_ALIGNMENT.CENTER)
    ts.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT)
    p._p.append(run_el(tab=True))
    math = parse_xml(
        f'<m:oMath {nsdecls("m")}><m:r><m:t>{formula}</m:t></m:r></m:oMath>'
    )
    p._p.append(math)
    p._p.append(run_el(tab=True))
    p._p.append(run_el(number))
    return p


def build_captions():
    doc = Document()
    setup_styles(doc)
    cover_and_front(doc)
    body_section(doc)
    content = CONTENT
    i1 = content.index(("h2", "1.2 研究意义"))
    add_content(doc, content[:i1])
    ref_para(doc, "论文格式检查的一般流程如", "图 1-1", "_Ref170000001", "所示，先解析文件，再逐项比对规则。")
    figure(doc, 1, 1, "格式检查的一般流程", "_Ref170000001")
    doc.add_paragraph("实验所用的软硬件环境见表 1-1。表中列出了操作系统、运行环境和测试浏览器的版本。")
    small_table(doc, 1, 1, "实验环境配置", [["项目", "配置"], ["操作系统", "Ubuntu 22.04"], ["运行环境", "Node.js 20"], ["浏览器", "Chrome 120"]])
    doc.add_paragraph("在表格之后继续正文叙述，说明实验环境对结果的影响很小，可以忽略不计。")
    i2 = content.index(("h2", "2.2 本章小结"))
    add_content(doc, content[i1:i2])
    doc.add_paragraph("注意力的计算公式如下：")
    # 较长的居中公式行（超过 40 个字符）
    equation(doc, "Attention(Q, K, V) = softmax(QKᵀ/√d)V    (2-1)")
    p = doc.add_paragraph("其中 Q、K、V 分别表示查询、键和值矩阵，d 表示键向量的维度。缩放因子用于防止点积过大。")
    assert p.text
    tab_equation(doc, "PE(pos,2i)=sin(pos/10000^(2i/d))", "(2-2)")
    doc.add_paragraph("式（2-2）给出了正弦位置编码的定义，其中 pos 表示位置，i 表示维度序号。")
    figure(doc, 2, 1, "Transformer 编码器结构", "_Ref170000002")
    figure(doc, 2, 2, "多头注意力示意", "_Ref170000003")
    small_table(doc, 2, 1, "不同位置编码方式的比较", [["方式", "可外推", "参数量"], ["正弦编码", "是", "0"], ["可学习编码", "否", "n×d"]])
    i3 = content.index(("h2", "3.2 模块划分"))
    add_content(doc, content[i2:i3])
    doc.add_paragraph("训练时使用交叉熵损失函数，其定义如下：")
    equation(doc, "L = −Σ y log ŷ    (3-1)")  # 短的居中公式行
    add_content(doc, content[i3:])
    back_matter(doc)
    return doc


# ---------- 4. 自动编号标题 + 混合列表 ----------

def numbering_xml_parts():
    def lvl(ilvl, fmt, text, style=None, left=0, hanging=0, legal=False, suff="space"):
        return (
            f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
            + (f'<w:pStyle w:val="{style}"/>' if style else "")
            + ('<w:isLgl/>' if legal else "")
            + f'<w:suff w:val="{suff}"/><w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
            + f'<w:pPr><w:ind w:left="{left}" w:hanging="{hanging}"/></w:pPr></w:lvl>'
        )

    heading = (
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90"><w:multiLevelType w:val="multilevel"/>'
        + lvl(0, "chineseCounting", "第%1章", "Heading1")
        + lvl(1, "decimal", "%1.%2", "Heading2", legal=True)
        + lvl(2, "decimal", "%1.%2.%3", "Heading3", legal=True)
        + "</w:abstractNum>"
    )
    bullets = (
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="91"><w:multiLevelType w:val="hybridMultilevel"/>'
        + lvl(0, "decimal", "%1.", left=840, hanging=360, suff="tab")
        + "</w:abstractNum>"
    )
    refs = (
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="92"><w:multiLevelType w:val="singleLevel"/>'
        + lvl(0, "decimal", "[%1]", left=420, hanging=420, suff="tab")
        + "</w:abstractNum>"
    )
    nums = [
        f'<w:num {nsdecls("w")} w:numId="{n}"><w:abstractNumId w:val="{a}"/></w:num>'
        for n, a in ((90, 90), (91, 91), (92, 92))
    ]
    return [parse_xml(x) for x in (heading, bullets, refs)], [parse_xml(x) for x in nums]


def num_pr(ppr_owner, num_id, ilvl=None):
    ppr = ppr_owner.get_or_add_pPr()
    np_ = OxmlElement("w:numPr")
    if ilvl is not None:
        np_.append(el("w:ilvl", val=ilvl))
    np_.append(el("w:numId", val=num_id))
    # numPr 要排在 pPr 里靠前的位置（pStyle、keepNext、keepLines、pageBreakBefore、framePr、widowControl 之后）
    anchor = None
    for tag in ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl"):
        found = ppr.find(qn(tag))
        if found is not None:
            anchor = found
    if anchor is None:
        ppr.insert(0, np_)
    else:
        anchor.addnext(np_)


def build_mixed_numbering():
    doc = Document()
    setup_styles(doc)
    numbering = doc.part.numbering_part.element
    abstracts, nums = numbering_xml_parts()
    first_num = numbering.find(qn("w:num"))
    for a in abstracts:
        if first_num is not None:
            first_num.addprevious(a)
        else:
            numbering.append(a)
    for n in nums:
        numbering.append(n)
    for level in (1, 2, 3):
        num_pr(doc.styles[f"Heading {level}"].element, 90, None if level == 1 else level - 1)

    def no_number(h):
        num_pr(h._p, 0)

    # 前置部分的标题套"标题 1"但取消编号
    orig_add_heading = doc.add_heading

    def add_heading(text, level=1):
        h = orig_add_heading(text, level)
        if re.sub(r"\s", "", text) in ("摘要", "ABSTRACT", "目录", "参考文献", "致谢"):
            no_number(h)
        return h

    doc.add_heading = add_heading
    cover_and_front(doc, toc=lambda d: [
        d.add_paragraph(t).paragraph_format.__setattr__("first_line_indent", Pt(0))
        for t in ("第一章 绪论 ........ 1", "1.1 研究背景 ........ 1", "第二章 研究方法 ........ 2")
    ])
    body_section(doc)

    strip = lambda t: re.sub(r"^(第[一二三四五六七八九十]+章|\d+(\.\d+)*)\s*", "", t)  # noqa: E731
    i = CONTENT.index(("h1", "第二章 相关技术"))
    add_content(doc, [(k, strip(t) if k != "p" else t) for k, t in CONTENT[:i]])

    doc.add_heading("研究方法", 1)
    doc.add_heading("研究步骤", 2)
    doc.add_paragraph("本研究按以下步骤进行，每一步都有明确的输入和输出：")
    for t in ("（1）数据收集：从公开渠道收集近五年的本科毕业论文模板，共计 120 份。",
              "（2）规则整理：逐条整理各校的格式要求，形成结构化的规则档案。",
              "（3）工具实现：在浏览器中实现解析和检查逻辑，不上传任何文件。"):
        doc.add_paragraph(t)
    doc.add_heading("评价指标", 2)
    doc.add_paragraph("评价时主要关注以下指标，它们分别从不同角度反映检查工具的效果：")
    for t in ("误报率，即合格的论文被报告为有问题的比例；",
              "漏报率，即有问题的论文没有被检出的比例；",
              "运行时间，即从上传文件到生成报告所需的时间。"):
        p = doc.add_paragraph(t)
        num_pr(p._p, 91, 0)
        p.paragraph_format.first_line_indent = None
    doc.add_paragraph("1）首先计算误报率；2）然后计算漏报率；3）最后统计运行时间。三个指标综合起来评价工具。")
    doc.add_heading("指标计算", 3)
    doc.add_paragraph("第一，误报率按样例数计算；第二，漏报率按错误点计算，这样可以避免一个样例里的多处错误被合并。")
    rest = CONTENT[CONTENT.index(("h1", "第三章 系统设计")):]
    add_content(doc, [(k, strip(t) if k != "p" else t) for k, t in rest])
    # 前面"第二章 相关技术"被替换掉了，引用 [5] 补在这里
    doc.add_paragraph("系统的核心模型采用 Transformer 结构[5]，这里不再赘述其原理。")

    def ref_numbered(p):
        num_pr(p._p, 92, 0)
        p.paragraph_format.first_line_indent = None

    back_matter(doc, numbered=ref_numbered)
    return doc


# ---------- 5. WPS 风格 ----------

def wps_rpr(cn="宋体", en="Times New Roman", half_pt=24, bold=False):
    b = "<w:b/><w:bCs/>" if bold else ""
    return parse_xml(
        f'<w:rPr {nsdecls("w")}><w:rFonts w:hint="eastAsia" w:ascii="{en}" w:hAnsi="{en}" '
        f'w:eastAsia="{cn}" w:cs="{en}"/>{b}<w:sz w:val="{half_pt}"/><w:szCs w:val="{half_pt}"/>'
        f'<w:lang w:val="en-US" w:eastAsia="zh-CN"/></w:rPr>'
    )


def wps_ppr(style=None, jc="both", first_chars=200, line=360, before=0, after=0, outline=None):
    s = f'<w:pStyle w:val="{style}"/>' if style else ""
    ol = f'<w:outlineLvl w:val="{outline}"/>' if outline is not None else ""
    ind = (
        f'<w:ind w:firstLine="{first_chars * 12 // 5}" w:firstLineChars="{first_chars}"/>'
        if first_chars else '<w:ind w:left="0" w:leftChars="0" w:firstLine="0" w:firstLineChars="0"/>'
    )
    return parse_xml(
        f'<w:pPr {nsdecls("w")}>{s}<w:keepNext w:val="0"/><w:keepLines w:val="0"/><w:widowControl/>'
        f'<w:snapToGrid w:val="0"/><w:spacing w:before="{before}" w:beforeLines="0" w:after="{after}" '
        f'w:afterLines="0" w:line="{line}" w:lineRule="auto"/>{ind}<w:jc w:val="{jc}"/>{ol}'
        f'<w:rPr><w:rFonts w:hint="eastAsia"/></w:rPr></w:pPr>'
    )


HEAD_FMT = {1: (32, "center"), 2: (28, "left"), 3: (24, "left")}


def wps_para(doc, text, kind="p"):
    body = doc.element.body
    p = OxmlElement("w:p")
    if kind in ("h1", "h2", "h3"):
        lvl = int(kind[1])
        size, jc = HEAD_FMT[lvl]
        p.append(wps_ppr(f"Heading{lvl}", jc=jc, first_chars=0, before=240, after=240, outline=lvl - 1))
        rpr = lambda: wps_rpr("黑体", "Times New Roman", size, bold=True)  # noqa: E731
    elif kind == "ref":
        p.append(wps_ppr(first_chars=0))
        rpr = lambda: wps_rpr(half_pt=21)  # noqa: E731
    elif kind == "center":
        p.append(wps_ppr(jc="center", first_chars=0))
        rpr = lambda: wps_rpr()  # noqa: E731
    else:
        p.append(wps_ppr())
        rpr = lambda: wps_rpr()  # noqa: E731
    # WPS 常把一段拆成很多 run（中文、西文、标点分开），每个 run 都带完整的直接格式
    for piece in re.findall(r"[A-Za-z0-9 .\-()]+|[^A-Za-z0-9 .\-()]+", text):
        p.append(run_el(piece, rpr=rpr()))
    body.insert(len(body) - 1, p)  # 放在最后的 sectPr 之前
    return p


def wps_section_break(doc, p_el, sect_xml):
    ppr = p_el.find(qn("w:pPr"))
    ppr.append(parse_xml(sect_xml))


def wps_sectpr(refs="", pgnum="", ns=True):
    return (
        f'<w:sectPr {nsdecls("w", "r")}>{refs}<w:pgSz w:w="11906" w:h="16838"/>'
        f'<w:pgMar w:top="1417" w:right="1417" w:bottom="1417" w:left="1701" w:header="851" w:footer="992" w:gutter="0"/>'
        f'{pgnum}<w:cols w:space="425" w:num="1"/><w:docGrid w:type="lines" w:linePitch="312" w:charSpace="0"/></w:sectPr>'
    )


def build_wps():
    """先用 python-docx 建好页眉页脚部件，再用直接格式写段落，最后改写 zip 里的样式 id 和 app.xml。"""
    doc = Document()
    core = doc.core_properties
    core.author = core.last_modified_by = ""
    core.created = core.modified = datetime.datetime(2026, 3, 1)

    # 样式里不放字体字号（WPS 把格式都写在 run 上）；标题样式只有大纲级别
    for level in (1, 2, 3):
        st = doc.styles[f"Heading {level}"]
        rpr = st.element.find(qn("w:rPr"))
        if rpr is not None:
            st.element.remove(rpr)

    # 三节：封面 / 前置（罗马）/ 正文（阿拉伯）。页眉页脚部件借 python-docx 创建
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_section(WD_SECTION.NEW_PAGE)
    # 注意 add_section 返回的对象指向末尾的 sectPr，再加一节后就变成了下一节，所以重新取
    for sec in list(doc.sections)[1:]:
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        sec.header.paragraphs[0].text = HEADER
        sec.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        wps_footer(sec.footer)
    # 拿到 python-docx 生成的三个 sectPr 里的页眉页脚引用，然后清空 body 重写
    body = doc.element.body
    sectprs = [p.find(qn("w:pPr")).find(qn("w:sectPr")) for p in body.findall(qn("w:p")) if p.find(qn("w:pPr")) is not None and p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None]
    final = body.find(qn("w:sectPr"))
    refs = []
    for sp in sectprs[1:] + [final]:
        refs.append("".join(
            f'<w:{c.tag.split("}")[1]} w:type="{c.get(qn("w:type"))}" r:id="{c.get(qn("r:id"))}"/>'
            for c in sp if c.tag in (qn("w:headerReference"), qn("w:footerReference"))
        ))
    for c in list(body):
        if c.tag != qn("w:sectPr"):
            body.remove(c)

    # 页边距：上下 2.5 cm（1417）、左 3.0 cm（1701）、右 2.5 cm（1417）
    wps_para(doc, "某某大学本科毕业论文", "center")
    p = wps_para(doc, "论文题目：示例论文", "center")
    wps_section_break(doc, p, wps_sectpr())

    wps_para(doc, "摘  要", "h1")
    wps_para(doc, ABSTRACT_CN)
    wps_para(doc, "ABSTRACT", "h1")
    wps_para(doc, ABSTRACT_EN)
    wps_para(doc, "目  录", "h1")
    last = None
    for k, t in CONTENT:
        if k in ("h1", "h2"):
            last = wps_para(doc, f"{t}　　1", "center")
    wps_section_break(doc, last, wps_sectpr(refs[0], '<w:pgNumType w:fmt="upperRoman" w:start="1"/>'))

    for k, t in CONTENT:
        wps_para(doc, t, k)
    wps_para(doc, "参考文献", "h1")
    for i, r in enumerate(ALL_REFS, 1):
        wps_para(doc, f"[{i}] {r}", "ref")
    wps_para(doc, "致  谢", "h1")
    wps_para(doc, "感谢指导教师在论文写作过程中给予的帮助，也感谢同学们的支持。")

    body.remove(final)
    body.append(parse_xml(wps_sectpr(refs[1], '<w:pgNumType w:fmt="decimal" w:start="1"/>')))
    return doc


def wps_footer(footer):
    """WPS 的页码：页脚段落里放一个文本框（mc:AlternateContent + VML 回退），框里是居中的 PAGE 域。"""
    ftr = footer._element
    for p in ftr.findall(qn("w:p")):
        ftr.remove(p)

    def inner():
        return (
            '<w:txbxContent><w:p><w:pPr><w:pStyle w:val="Footer"/><w:jc w:val="center"/></w:pPr>'
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> PAGE  \\* MERGEFORMAT </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>1</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r>'
            "</w:p></w:txbxContent>"
        )

    xml = (
        f'<w:p {nsdecls("w", "wp", "a", "pic")} {EXTRA_NS}>'
        '<w:pPr><w:pStyle w:val="Footer"/></w:pPr><w:r><mc:AlternateContent>'
        '<mc:Choice Requires="wps"><w:drawing><wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" '
        'relativeHeight="251659264" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="margin"><wp:align>center</wp:align></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        '<wp:extent cx="1828800" cy="1828800"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        '<wp:docPr id="1" name="文本框 1"/><wp:cNvGraphicFramePr/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1828800" cy="1828800"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln w="6350"><a:noFill/></a:ln></wps:spPr>'
        f'<wps:txbx>{inner()}</wps:txbx>'
        '<wps:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" upright="1"><a:spAutoFit/></wps:bodyPr>'
        '</wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>'
        '<mc:Fallback><w:pict><v:shape id="文本框 1" o:spid="_x0000_s1025" type="#_x0000_t202" '
        'style="position:absolute;margin-top:0pt;height:144pt;width:144pt;mso-position-horizontal:center;'
        'mso-position-horizontal-relative:margin;mso-wrap-style:none;" filled="f" stroked="f" coordsize="21600,21600">'
        f'<v:textbox inset="0mm,0mm,0mm,0mm" style="mso-fit-shape-to-text:t;">{inner()}</v:textbox>'
        '</v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>'
    )
    ftr.append(parse_xml(xml))


EXTRA_NS = (
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
)

WPS_IDS = {"Normal": "1", "Heading1": "2", "Heading2": "3", "Heading3": "4", "Header": "5", "Footer": "6",
           "DefaultParagraphFont": "7", "TableNormal": "8"}

WPS_APP = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><Template>Normal</Template><Pages>6</Pages><Words>1200</Words><Characters>1500</Characters><Lines>40</Lines><Paragraphs>30</Paragraphs><TotalTime>0</TotalTime><ScaleCrop>false</ScaleCrop><LinksUpToDate>false</LinksUpToDate><CharactersWithSpaces>1600</CharactersWithSpaces><Application>WPS Office_12.1.0.17827_F1E327BC-269C-435d-A252-1D8C7A1D1A3F</Application><DocSecurity>0</DocSecurity></Properties>"""


def wps_postprocess(path: Path):
    """样式 id 改成数字（名字改成小写的 heading 1 这类），app.xml 写 WPS。"""
    src = zipfile.ZipFile(path)
    items = {n: src.read(n) for n in src.namelist()}
    src.close()
    for name in list(items):
        if not (name.startswith("word/") and name.endswith(".xml")):
            continue
        xml = items[name].decode("utf-8")
        for old, new in WPS_IDS.items():
            xml = re.sub(rf'(w:styleId=")({old})(")', rf"\g<1>{new}\g<3>", xml)
            xml = re.sub(rf'(<w:(?:pStyle|rStyle|tblStyle|basedOn|next|link) w:val=")({old})(")', rf"\g<1>{new}\g<3>", xml)
        if name == "word/styles.xml":
            xml = re.sub(r'(<w:name w:val=")Heading (\d)(")', r"\1heading \2\3", xml)
        items[name] = xml.encode("utf-8")
    items["docProps/app.xml"] = WPS_APP.encode("utf-8")
    items.pop("word/stylesWithEffects.xml", None)
    ct = items["[Content_Types].xml"].decode()
    ct = re.sub(r'<Override PartName="/word/stylesWithEffects.xml"[^>]*/>', "", ct)
    items["[Content_Types].xml"] = ct.encode()
    rels = items["word/_rels/document.xml.rels"].decode()
    rels = re.sub(r'<Relationship [^>]*Target="stylesWithEffects.xml"[^>]*/>', "", rels)
    items["word/_rels/document.xml.rels"] = rels.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in items.items():
            z.writestr(n, data)


# ---------- 6. LibreOffice 往返 ----------

def libreoffice_roundtrip(src: Path, dst: Path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("没有找到 soffice，跳过", dst.name)
        return
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp) / "profile"
        env_arg = f"-env:UserInstallation=file://{profile}"
        subprocess.run([soffice, env_arg, "--headless", "--convert-to", "odt", "--outdir", tmp, str(src)], check=True, capture_output=True)
        odt = Path(tmp) / (src.stem + ".odt")
        outdir = Path(tmp) / "back"
        outdir.mkdir()
        subprocess.run([soffice, env_arg, "--headless", "--convert-to", "docx:MS Word 2007 XML", "--outdir", str(outdir), str(odt)], check=True, capture_output=True)
        shutil.copy(outdir / (src.stem + ".docx"), dst)


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    build_toc_field().save(OUT / "toc-field.docx")
    build_landscape().save(OUT / "landscape.docx")
    build_captions().save(OUT / "captions.docx")
    build_mixed_numbering().save(OUT / "mixed-numbering.docx")
    build_wps().save(OUT / "wps-like.docx")
    wps_postprocess(OUT / "wps-like.docx")
    libreoffice_roundtrip(OUT / "captions.docx", OUT / "captions-libreoffice.docx")
    print("wrote", ", ".join(sorted(p.name for p in OUT.glob("*.docx"))))
